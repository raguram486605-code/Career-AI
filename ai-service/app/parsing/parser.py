"""Resume parsing: PDF/DOCX -> text -> structured data."""
import io
import re
from ..skills.extractor import get_extractor

SECTION_HEADERS = {
    "skills": r"(technical\s+)?skills|technologies|tech\s+stack|core\s+competenc(y|ies)",
    "projects": r"(academic\s+|personal\s+)?projects?",
    "experience": r"(work\s+|professional\s+)?experience|internships?|employment",
    "education": r"education|academic\s+(background|qualification)s?",
    "certifications": r"certifications?|courses|achievements|awards",
    "summary": r"(career\s+)?(summary|objective)|profile|about\s+me",
}


def extract_text(file_bytes: bytes, filename: str) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        import pdfplumber
        parts = []
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                parts.append(page.extract_text() or "")
        return "\n".join(parts)
    if name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(file_bytes))
        lines = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                lines.append(" | ".join(c.text for c in row.cells))
        return "\n".join(lines)
    if name.endswith(".txt"):
        return file_bytes.decode("utf-8", errors="ignore")
    raise ValueError("Unsupported file type. Upload PDF, DOCX or TXT.")


def split_sections(text: str) -> dict:
    """Heuristic section splitter: a short line that matches a header keyword starts a section."""
    sections, current = {"header": []}, "header"
    for line in text.splitlines():
        clean = line.strip().strip(":-—_ ").strip()
        key = None
        if 0 < len(clean) <= 40:
            for k, pat in SECTION_HEADERS.items():
                if re.fullmatch(pat, clean, re.IGNORECASE):
                    key = k
                    break
        if key:
            current = key
            sections.setdefault(current, [])
        else:
            sections.setdefault(current, []).append(line)
    return {k: "\n".join(v).strip() for k, v in sections.items() if "\n".join(v).strip()}


def _email(t):
    m = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", t)
    return m.group(0) if m else None


def _phone(t):
    m = re.search(r"(\+?\d[\d\s\-()]{8,14}\d)", t)
    return m.group(1).strip() if m else None


def _cgpa(t):
    m = re.search(r"(?:cgpa|gpa)\s*[:\-]?\s*(\d(?:\.\d{1,2})?)\s*(?:/\s*10)?", t, re.IGNORECASE)
    if m:
        v = float(m.group(1))
        return v if v <= 10 else None
    return None


def parse_resume(text: str) -> dict:
    ex = get_extractor()
    sections = split_sections(text)
    all_counts = ex.extract(text)
    project_text = "\n".join(sections.get(k, "") for k in ("projects", "experience"))
    project_skills = ex.skills(project_text)
    declared = ex.skills(sections.get("skills", ""))
    # Skills mentioned only in the free text still count, ranked by evidence
    ranked = sorted(all_counts, key=lambda k: (-all_counts[k], k))
    return {
        "email": _email(text),
        "phone": _phone(text),
        "cgpa": _cgpa(text),
        "skills": ranked,
        "declaredSkills": declared,
        "projectSkills": project_skills,
        "skillEvidence": all_counts,
        "sections": sections,
        "projectsText": project_text,
        "rawText": text,
        "wordCount": len(text.split()),
        "hasGithub": bool(re.search(r"github\.com", text, re.IGNORECASE)),
        "quality": resume_quality(text, sections, ranked),
    }


def resume_quality(text: str, sections: dict, skills: list) -> dict:
    """Simple, explainable resume-quality checklist (0-100)."""
    checks = {
        "Has a skills section": "skills" in sections,
        "Has projects": "projects" in sections,
        "Has education": "education" in sections,
        "Has experience / internships": "experience" in sections,
        "Contact email present": bool(_email(text)),
        "Mentions GitHub / portfolio": bool(re.search(r"github\.com|portfolio|linkedin\.com", text, re.I)),
        "Uses numbers / impact (e.g. 40%, 1000 users)": bool(re.search(r"\d+\s?(%|\+|users|ms|x)\b", text)),
        "Reasonable length (150-900 words)": 150 <= len(text.split()) <= 900,
        "At least 6 recognised skills": len(skills) >= 6,
    }
    score = round(100 * sum(checks.values()) / len(checks))
    tips = [f"Add / improve: {k}" for k, ok in checks.items() if not ok]
    return {"score": score, "checks": checks, "suggestions": tips}
